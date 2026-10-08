"""
Script Template Service providing sample screenplay templates (JSON, TXT, SRT, CSV, Excel XLSX)
and battle-tested AI prompt templates for ChatGPT, Claude AI, DeepSeek, and Gemini.
Pure Python application service adhering to Clean Architecture.
"""

import json
from pathlib import Path
from typing import Dict, List, Tuple, Union


class ScriptTemplateService:
    """Manages built-in script templates, exporters, and AI prompt suggestions."""

    # ─────────────────────────────────────────────────────────────
    # Sample Script Data
    # ─────────────────────────────────────────────────────────────

    SAMPLE_TITLE = "Khám Phá Thiên Nhiên Kỳ Vĩ Việt Nam"

    SAMPLE_SCENES = [
        {
            "id": 1,
            "time_start": "00:00:00",
            "time_end": "00:00:05",
            "duration_seconds": 5.0,
            "dialogue_es": "Chào mừng các bạn đến với hành trình khám phá những danh lam thắng cảnh hùng vĩ nhất Việt Nam.",
            "primary_keywords": ["vietnam nature", "vietnam landscape", "mountains"],
            "secondary_keywords": ["drone view", "travel adventure"]
        },
        {
            "id": 2,
            "time_start": "00:00:05",
            "time_end": "00:00:10",
            "duration_seconds": 5.0,
            "dialogue_es": "Từ những dãy núi đá vôi trập trùng tại vịnh Hạ Long cho đến ruộng bậc thang óng ả Mù Cang Chải.",
            "primary_keywords": ["ha long bay", "mu cang chai", "rice terraces"],
            "secondary_keywords": ["golden sunset", "aerial vietnam"]
        },
        {
            "id": 3,
            "time_start": "00:00:10",
            "time_end": "00:00:15",
            "duration_seconds": 5.0,
            "dialogue_es": "Mỗi bước chân là một bức tranh thiên nhiên tuyệt mỹ mà tạo hóa đã ban tặng.",
            "primary_keywords": ["vietnam coastline", "tropical forest", "nature beauty"],
            "secondary_keywords": ["cinematic landscape", "wanderlust"]
        }
    ]

    # ─────────────────────────────────────────────────────────────
    # AI Prompt Templates
    # ─────────────────────────────────────────────────────────────

    PROMPTS: Dict[str, Dict[str, str]] = {
        "json": {
            "title": "Kịch bản JSON chuẩn (Claude / ChatGPT / Gemini)",
            "description": "Tạo kịch bản cấu trúc JSON đầy đủ phân cảnh, từ khóa stock tiếng Anh và thời lượng chuẩn xác.",
            "content": """Bạn là chuyên gia biên kịch video và đạo diễn hình ảnh B-roll chuyên nghiệp.
Hãy viết một kịch bản phân đoạn video theo chủ đề: [CHỦ ĐỀ CỦA BẠN - Ví dụ: Khám Phá Công Nghệ AI 2026]
Thời lượng dự kiến: [Ví dụ: 60 giây, khoảng 10-12 cảnh]

Yêu cầu xuất ra DUY NHẤT một khối mã JSON hợp lệ theo đúng cấu trúc sau (không kèm lời chào hay giải thích ngoài mã):
```json
{
  "title": "[Tên video]",
  "total_scenes": 3,
  "scenes": [
    {
      "id": 1,
      "time_start": "00:00:00",
      "time_end": "00:00:05",
      "duration_seconds": 5.0,
      "dialogue_es": "Câu thoại hoặc lời bình của cảnh này (bằng tiếng Việt tự nhiên, hấp dẫn).",
      "primary_keywords": ["từ khóa tiếng Anh 1", "từ khóa tiếng Anh 2", "từ khóa 3"],
      "secondary_keywords": ["từ khóa bổ trợ 1", "từ khóa bổ trợ 2"]
    }
  ]
}
```

Lưu ý quan trọng:
1. 'primary_keywords' và 'secondary_keywords' PHẢI viết bằng TIẾNG ANH mô tả hình ảnh trực quan (ví dụ: 'futuristic city', 'coding on laptop', 'drone aerial sunset') để AutoStock Studio tự động tìm và tải stock video từ Pexels, Pixabay, Wikimedia, Coverr chính xác nhất.
2. 'duration_seconds' nên từ 4.0 đến 6.0 giây mỗi cảnh.
3. 'dialogue_es' là câu thoại tiếng Việt sẽ được chuyển thành giọng đọc AI bản địa."""
        },

        "excel": {
            "title": "Kịch bản Bảng Tính Excel / CSV",
            "description": "Tạo kịch bản dạng bảng 4 cột (Cảnh, Lời thoại, Từ khóa stock, Thời lượng) dễ chỉnh sửa trong Excel hoặc Google Sheets.",
            "content": """Bạn là trợ lý biên kịch video. Hãy soạn kịch bản video theo chủ đề: [CHỦ ĐỀ CỦA BẠN].
Số lượng cảnh: [Ví dụ: 8 cảnh].

Hãy xuất ra bảng dữ liệu theo định dạng CSV hoặc bảng Markdown có 4 cột sau:
- Cột 1: Cảnh (Số thứ tự 1, 2, 3...)
- Cột 2: Lời thoại (Câu thoại tiếng Việt tự nhiên)
- Cột 3: Từ khóa (3-5 từ khóa stock media tiếng Anh miêu tả cảnh quay, ngăn cách bởi dấu phẩy)
- Cột 4: Thời lượng (Thời lượng cảnh tính bằng giây, ví dụ: 4.5, 5, 6)

Ví dụ định dạng xuất ra:
Cảnh,Lời thoại,Từ khóa,Thời lượng
1,"Chào mừng các bạn đến với xu hướng công nghệ tương lai.","futuristic technology, neon city, modern",5
2,"Trí tuệ nhân tạo đang thay đổi cách chúng ta làm việc.","artificial intelligence, digital brain, futuristic lab",4.5
3,"Hãy cùng khám phá tiềm năng không giới hạn ngay hôm nay.","success, inspiration, creative workspace",5"""
        },

        "txt": {
            "title": "Kịch bản Văn Bản TXT Phân Cảnh",
            "description": "Tạo kịch bản văn bản ngắn gọn, mỗi dòng là một cảnh kèm thời lượng và từ khóa stock.",
            "content": """Hãy viết kịch bản video ngắn theo chủ đề: [CHỦ ĐỀ CỦA BẠN].
Mỗi dòng là một phân cảnh được định dạng chính xác theo cấu trúc:
Cảnh [STT] ([Thời lượng]s): [Lời thoại hoặc lời bình] | Từ khóa: [3-5 từ khóa stock tiếng Anh]

Ví dụ mẫu:
Cảnh 1 (5s): Chào mừng bạn đến với thế giới ẩm thực đường phố sôi động | Từ khóa: street food, night market, asian food
Cảnh 2 (4.5s): Mùi hương thơm lừng từ những món ăn truyền thống nóng hổi | Từ khóa: cooking fire, sizzling pan, chef cooking
Cảnh 3 (5s): Trải nghiệm những hương vị độc đáo làm say đắm lòng người | Từ khóa: happy eating, friends dining, delicious food"""
        },

        "srt": {
            "title": "Kịch bản Phụ Đề Video SRT",
            "description": "Tạo phụ đề video chuẩn SRT có đánh số thứ tự và timecode chuẩn xác từng mili-giây.",
            "content": """Hãy viết phụ đề kịch bản video dạng chuẩn SRT theo chủ đề: [CHỦ ĐỀ CỦA BẠN].
Thời lượng dự kiến: [Ví dụ: 30 giây].
Mỗi đoạn phụ đề cách nhau bởi một dòng trống, có chỉ số và mốc thời gian chuẩn dạng 00:00:00,000 --> 00:00:05,000.

Ví dụ mẫu:
1
00:00:00,000 --> 00:00:04,500
Khám phá những vùng đất mới luôn đem lại nguồn cảm hứng bất tận.

2
00:00:04,800 --> 00:00:09,500
Nơi thiên nhiên hoang sơ hòa quyện cùng bầu trời xanh ngút ngàn.

3
00:00:09,800 --> 00:00:15,000
Hãy xách ba lô lên và bắt đầu hành trình của riêng bạn ngay hôm nay."""
        }
    }

    # Aliases for convenience and compatibility
    for _item in PROMPTS.values():
        _item["desc"] = _item["description"]
        _item["prompt"] = _item["content"]

    # ─────────────────────────────────────────────────────────────
    # Format Generators
    # ─────────────────────────────────────────────────────────────

    @classmethod
    def get_sample_json(cls) -> str:
        """Returns pretty-printed sample JSON screenplay."""
        data = {
            "title": cls.SAMPLE_TITLE,
            "total_scenes": len(cls.SAMPLE_SCENES),
            "scenes": cls.SAMPLE_SCENES
        }
        return json.dumps(data, ensure_ascii=False, indent=2)

    @classmethod
    def get_sample_txt(cls) -> str:
        """Returns sample TXT screenplay."""
        lines = [
            f"# KỊCH BẢN MẪU: {cls.SAMPLE_TITLE.upper()}",
            "# Định dạng: Cảnh [STT] ([Thời lượng]s): [Lời thoại] | Từ khóa: [Từ khóa stock]",
            ""
        ]
        for s in cls.SAMPLE_SCENES:
            sid = s["id"]
            dur = s["duration_seconds"]
            dlg = s["dialogue_es"]
            kws = ", ".join(s["primary_keywords"] + s["secondary_keywords"][:1])
            lines.append(f"Cảnh {sid} ({dur}s): {dlg} | Từ khóa: {kws}")
        return "\n".join(lines) + "\n"

    @classmethod
    def get_sample_srt(cls) -> str:
        """Returns sample SRT subtitle file."""
        blocks = []
        for s in cls.SAMPLE_SCENES:
            sid = s["id"]
            start = s["time_start"]
            if "," not in start and "." not in start:
                start = f"{start},000"
            end = s["time_end"]
            if "," not in end and "." not in end:
                end = f"{end},000"
            dlg = s["dialogue_es"]
            blocks.append(f"{sid}\n{start} --> {end}\n{dlg}")
        return "\n\n".join(blocks) + "\n"

    @classmethod
    def get_sample_csv(cls) -> str:
        """Returns sample CSV table."""
        lines = ["Cảnh,Lời thoại,Từ khóa,Thời lượng"]
        for s in cls.SAMPLE_SCENES:
            sid = s["id"]
            dlg = f'"{s["dialogue_es"]}"'
            kws = f'"{", ".join(s["primary_keywords"])}"'
            dur = s["duration_seconds"]
            lines.append(f"{sid},{dlg},{kws},{dur}")
        return "\n".join(lines) + "\n"

    @classmethod
    def generate_sample_xlsx(cls, target_path: Path):
        """Generates sample Excel workbook at target_path."""
        try:
            import openpyxl  # type: ignore
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side  # type: ignore

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Kịch Bản Mẫu"

            # Headers
            headers = ["Cảnh", "Lời Thoại", "Từ Khóa Stock Media", "Thời Lượng (giây)"]
            ws.append(headers)

            # Header styling
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="38BDF8")
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = align_center

            # Data rows
            thin_border = Border(
                left=Side(style='thin', color='CBD5E1'),
                right=Side(style='thin', color='CBD5E1'),
                top=Side(style='thin', color='CBD5E1'),
                bottom=Side(style='thin', color='CBD5E1')
            )

            for s in cls.SAMPLE_SCENES:
                row_data = [
                    s["id"],
                    s["dialogue_es"],
                    ", ".join(s["primary_keywords"] + s["secondary_keywords"][:1]),
                    s["duration_seconds"]
                ]
                ws.append(row_data)

            # Styling data cells
            for row in ws.iter_rows(min_row=2, max_row=len(cls.SAMPLE_SCENES) + 1):
                for col_idx, cell in enumerate(row):
                    cell.border = thin_border
                    cell.font = Font(name="Segoe UI", size=10)
                    if col_idx in (0, 3):
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left

            # Column widths
            ws.column_dimensions['A'].width = 10
            ws.column_dimensions['B'].width = 65
            ws.column_dimensions['C'].width = 38
            ws.column_dimensions['D'].width = 18

            wb.save(target_path)

        except Exception:
            # Fallback to writing CSV if openpyxl fails
            csv_path = target_path.with_suffix(".csv")
            csv_path.write_text(cls.get_sample_csv(), encoding="utf-8-sig")

    @classmethod
    def get_prompt(cls, format_key: str = "json") -> str:
        """Returns the AI prompt content string for the given format key."""
        item = cls.PROMPTS.get(format_key.lower()) or cls.PROMPTS.get("json", {})
        return item.get("content", "")

    @classmethod
    def get_prompt_info(cls, format_key: str = "json") -> Dict[str, str]:
        """Returns the full prompt metadata dict for the given format key."""
        return cls.PROMPTS.get(format_key.lower()) or cls.PROMPTS.get("json", {})

    @classmethod
    def export_all_templates(cls, target_dir: Union[Path, str]) -> List[Path]:
        """
        Exports all 5 sample template files plus a README instruction into the target folder.
        Returns the list of created file paths.
        """
        target_path = Path(target_dir)
        target_path.mkdir(parents=True, exist_ok=True)
        created_files = []

        # 1. JSON
        json_path = target_path / "kich_ban_mau.json"
        json_path.write_text(cls.get_sample_json(), encoding="utf-8")
        created_files.append(json_path)

        # 2. TXT
        txt_path = target_path / "kich_ban_mau.txt"
        txt_path.write_text(cls.get_sample_txt(), encoding="utf-8")
        created_files.append(txt_path)

        # 3. SRT
        srt_path = target_path / "kich_ban_mau.srt"
        srt_path.write_text(cls.get_sample_srt(), encoding="utf-8")
        created_files.append(srt_path)

        # 4. CSV
        csv_path = target_path / "kich_ban_mau.csv"
        csv_path.write_text(cls.get_sample_csv(), encoding="utf-8-sig")
        created_files.append(csv_path)

        # 5. Excel XLSX
        xlsx_path = target_path / "kich_ban_mau.xlsx"
        try:
            cls.generate_sample_xlsx(xlsx_path)
            if xlsx_path.exists():
                created_files.append(xlsx_path)
        except Exception:
            pass

        # 6. Readme instruction guide
        readme_path = target_path / "README_HUONG_DAN.txt"
        readme_content = (
            "=================================================================\n"
            "   HƯỚNG DẪN NẠP KỊCH BẢN VÀO AUTOSTOCK STUDIO PRO\n"
            "=================================================================\n\n"
            "Thư mục này chứa các tệp mẫu kịch bản chuẩn cho mọi định dạng:\n\n"
            "1. kich_ban_mau.xlsx : Tệp bảng tính Excel có định dạng cột sẵn.\n"
            "   - Cột A (scene_id)    : Số thứ tự cảnh (1, 2, 3...)\n"
            "   - Cột B (description) : Lời thoại hoặc mô tả phân cảnh\n"
            "   - Cột C (keywords)    : Từ khóa stock tiếng Anh (ngăn cách bằng dấu phẩy)\n"
            "   - Cột D (duration)    : Thời lượng cảnh tính bằng giây (ví dụ: 4.5, 5, 6)\n\n"
            "2. kich_ban_mau.csv  : Bảng tính CSV (mở được bằng Excel, Google Sheets, Notepad).\n\n"
            "3. kich_ban_mau.txt  : Văn bản phân cảnh tự do.\n"
            "   - Cú pháp: Cảnh 1 (5s): Lời thoại | Từ khóa: keyword1, keyword2\n\n"
            "4. kich_ban_mau.srt  : Phụ đề video chuẩn SubRip (.srt).\n\n"
            "5. kich_ban_mau.json : Cấu trúc JSON chi tiết, hỗ trợ đa trường dữ liệu.\n\n"
            "CÁCH SỬ DỤNG:\n"
            "- Điền kịch bản của bạn vào một trong các tệp trên.\n"
            "- Mở Stock Studio Pro -> Nhấn nút 'Nạp Kịch Bản' (Ctrl+O) -> Chọn tab 'Từ Tệp' và chọn file.\n"
            "- Hoặc kéo thả trực tiếp tệp kịch bản vào cửa sổ ứng dụng!\n"
        )
        readme_path.write_text(readme_content, encoding="utf-8")
        created_files.append(readme_path)

        return created_files

